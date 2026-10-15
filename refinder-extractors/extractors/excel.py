"""XLSX: nonoverlapping 3×3 tiles, with column-major traversal and coordinates."""
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter
from ..normalize import normalize_text

TILE_SIZE = 3


def extract_xlsx(path: Path) -> list[str]:
    # data_only skips the formula expressions; values are Excel's cached results.
    # A formula without a cached result may appear empty. Formulas are v2.
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    chunks = []
    try:
        for sheet in workbook.worksheets:
            # Stream each horizontal band once (read_only iterators can be costly
            # to restart for every separate tile).
            for top in range(1, sheet.max_row + 1, TILE_SIZE):
                rows = list(sheet.iter_rows(
                    min_row=top, max_row=min(top + TILE_SIZE - 1, sheet.max_row),
                    values_only=True,
                ))
                for left in range(1, sheet.max_column + 1, TILE_SIZE):
                    right = min(left + TILE_SIZE - 1, sheet.max_column)
                    bottom = top + len(rows) - 1
                    cells = []
                    for col in range(left, right + 1):
                        for offset, row in enumerate(rows):
                            value = row[col - 1] if col - 1 < len(row) else None
                            text = normalize_text(value)
                            if text:
                                address = f"{get_column_letter(col)}{top + offset}"
                                cells.append(f"{address}: {text}")
                    if cells:
                        bounds = (f"{get_column_letter(left)}{top}:"
                                  f"{get_column_letter(right)}{bottom}")
                        chunks.append(f"Sheet: {sheet.title}; Range: {bounds}\n" + "\n".join(cells))
    finally:
        workbook.close()
    return chunks
