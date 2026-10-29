"""PPTX: each slide is a single chunk, including shape text and table cells."""
from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from ..normalize import normalize_text


def _shape_text(shape) -> list[str]:
    parts = []
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for child in shape.shapes:
            parts.extend(_shape_text(child))
    if shape.has_text_frame:
        text = normalize_text(shape.text_frame.text)
        if text:
            parts.append(text)
    if shape.has_table:
        for row in shape.table.rows:
            values = [normalize_text(cell.text) for cell in row.cells]
            if any(values):
                parts.append(" | ".join(values))
    return parts


def extract_pptx(path: Path) -> list[str]:
    presentation = Presentation(path)
    chunks = []
    for slide_number, slide in enumerate(presentation.slides, 1):
        parts = []
        for shape in slide.shapes:
            parts.extend(_shape_text(shape))
        if parts:
            chunks.append(f"Slide {slide_number}\n" + "\n".join(parts))
    return chunks
