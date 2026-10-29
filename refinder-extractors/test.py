import json
import sqlite3
from pathlib import Path
from tempfile import TemporaryDirectory
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfWriter
from refinder.indexer import index_folders


def test_full_pipeline():
    with TemporaryDirectory() as tmp:
        base = Path(tmp)
        docs = base / 'docs'
        docs.mkdir()
        excluded = docs / 'private'
        excluded.mkdir()
        (docs / 'notes.txt').write_text('First paragraph.\n\nSecond paragraph.', encoding='utf-8')
        (excluded / 'secret.txt').write_text('secret', encoding='utf-8')
        (docs / 'data.csv').write_text('name,price\nPen,3\n', encoding='utf-8')

        word = Document()
        word.add_paragraph('Opening paragraph')
        word.add_table(rows=1, cols=2).rows[0].cells[0].text = 'Hello'
        word.save(docs / 'report.docx')

        wb = Workbook()
        ws = wb.active
        ws['A1'] = 'Heading'
        ws['A2'] = 'Vertical'
        ws['B1'] = 'Horizontal'
        ws['D4'] = 'Another tile'
        wb.save(docs / 'table.xlsx')

        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
        box.text = 'Slide body'
        deck.save(docs / 'slides.pptx')

        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        with (docs / 'empty.pdf').open('wb') as f:
            writer.write(f)

        db = base / 'refinder.db'
        result = index_folders([docs], [excluded], db)
        assert result == {'indexed': 5, 'empty': 1, 'error': 0, 'skipped': 0}, result
        with sqlite3.connect(db) as conn:
            found = {name: json.loads(payload) for name, payload in conn.execute(
                'SELECT filename,chunks_json FROM files')}
        assert 'secret.txt' not in found
        assert 'refinder.db' not in found
        assert len(found['notes.txt']) == 2
        assert found['slides.pptx'] == ['Slide 1\nSlide body']
        assert found['table.xlsx'][0].find('A2: Vertical') < found['table.xlsx'][0].find('B1: Horizontal')
        assert len(found['table.xlsx']) == 2
        assert found['empty.pdf'] == []
        assert index_folders([docs], [excluded], db)['skipped'] == 6
