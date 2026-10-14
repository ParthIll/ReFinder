"""Turn files on disk into plain text for embedding. Shared by the indexer and the evaluation harness."""

import re
from pathlib import Path

from docx import Document
from pypdf import PdfReader

# Anything not listed here is tried as UTF-8 text and skipped if it looks binary.
SKIP_SUFFIXES = {
    ".png", ".jpg", ".jpeg", ".gif", ".heic", ".webp", ".bmp", ".ico", ".mp3", ".wav", ".m4a", ".mp4", ".mov",
    ".zip", ".gz", ".tar", ".7z", ".dmg", ".exe", ".bin", ".so", ".dylib", ".pyc", ".db", ".sqlite",
}


class UnsupportedFile(Exception):
    pass


def _pdf(path: Path) -> str:
    return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)


def _docx(path: Path) -> str:
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _plain(path: Path) -> str:
    raw = path.read_bytes()
    if b"\x00" in raw[:8192]:
        raise UnsupportedFile("binary content")
    return raw.decode("utf-8", errors="ignore")


EXTRACTORS = {".pdf": _pdf, ".docx": _docx}


def extract_text(path: Path) -> str:
    """Returns the file's text (possibly empty, e.g. a scanned PDF). Raises UnsupportedFile for binaries."""
    suffix = path.suffix.lower()
    if suffix in SKIP_SUFFIXES or path.name.startswith("."):
        raise UnsupportedFile(f"skipped type {suffix or path.name}")
    text = EXTRACTORS.get(suffix, _plain)(path)
    # PDFs in particular pad lines with runs of spaces; they waste tokens in every chunk.
    text = re.sub(r"[ \t\r\f\v]+", " ", text)
    return re.sub(r"\n\s*\n\s*", "\n\n", text).strip()
